import json
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="D3C7B6"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'<w:right w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

# Load complete master dataset (Sets 1 to 16)
with open('all_sets_1_to_16.json', 'r', encoding='utf-8') as f:
    DATA = json.load(f)

def generate_kanji_document(subset_dict, filename="Kanji_Practice_Sheets_Sets_1_to_16.docx"):
    doc = Document()
    
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.4)
        s.bottom_margin = Inches(0.4)
        s.left_margin = Inches(0.5)
        s.right_margin = Inches(0.5)

    COL_WIDTHS = [Inches(1.8)] + [Inches(0.8)] * 7
    
    for set_idx, (set_name, kanji_list) in enumerate(subset_dict.items()):
        p_set = doc.add_paragraph()
        p_set.paragraph_format.space_before = Pt(6)
        p_set.paragraph_format.space_after = Pt(2)
        run_set = p_set.add_run(set_name)
        run_set.font.name = "Georgia"
        run_set.font.size = Pt(16)
        run_set.font.bold = True
        run_set.font.color.rgb = RGBColor(0x8C, 0x2A, 0x1E)

        for kanji, meaning, vocab_list in kanji_list:
            # Header Kanji Banner Table (1x1)
            header_tbl = doc.add_table(rows=1, cols=1)
            header_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            h_cell = header_tbl.rows[0].cells[0]
            h_cell.width = Inches(7.4)
            set_cell_background(h_cell, "F3EFE9")
            set_cell_margins(h_cell, top=40, bottom=40, left=120, right=120)
            
            tcPr = h_cell._element.get_or_add_tcPr()
            borders = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>'
                f'<w:left w:val="single" w:sz="36" w:space="0" w:color="8C2A1E"/>'
                f'<w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/>'
                f'</w:tcBorders>'
            )
            tcPr.append(borders)

            hp = h_cell.paragraphs[0]
            hp.paragraph_format.space_before = Pt(1)
            hp.paragraph_format.space_after = Pt(1)
            k_run = hp.add_run(kanji)
            k_run.font.name = "MS Mincho"
            k_run.font.size = Pt(22)
            k_run.font.bold = True

            m_p = h_cell.add_paragraph()
            m_p.paragraph_format.space_before = Pt(0)
            m_p.paragraph_format.space_after = Pt(1)
            m_run = m_p.add_run(meaning)
            m_run.font.name = "Georgia"
            m_run.font.size = Pt(9.5)
            m_run.font.italic = True
            m_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

            doc.add_paragraph().paragraph_format.space_after = Pt(2)

            # Practice Grid Table for this Kanji
            table = doc.add_table(rows=1, cols=8)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_table_borders(table, color="D3C7B6")

            # Header Row
            hdr_row = table.rows[0]
            hdr_titles = ["WORD", "1", "2", "3", "4", "5", "6", "7"]
            for i, cell in enumerate(hdr_row.cells):
                cell.width = COL_WIDTHS[i]
                set_cell_background(cell, "EDE5D8")
                set_cell_margins(cell, top=30, bottom=30, left=40, right=40)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                r = p.add_run(hdr_titles[i])
                r.font.name = "Arial"
                r.font.size = Pt(8)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x77, 0x77, 0x77)

            # Add 3 practice sub-rows per vocabulary word
            for w_kanji, w_hira, w_eng in vocab_list:
                r0 = table.add_row()
                r1 = table.add_row()
                r2 = table.add_row()

                for r in [r0, r1, r2]:
                    trPr = r._element.get_or_add_trPr()
                    trHeight = parse_xml(f'<w:trHeight {nsdecls("w")} w:val="850" w:hRule="atLeast"/>')
                    trPr.append(trHeight)
                    for i, c in enumerate(r.cells):
                        c.width = COL_WIDTHS[i]
                        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

                # Merge column 0 across r0, r1, r2
                info_cell = r0.cells[0]
                info_cell.merge(r1.cells[0])
                info_cell.merge(r2.cells[0])
                set_cell_margins(info_cell, top=40, bottom=40, left=80, right=40)
                info_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

                ip = info_cell.paragraphs[0]
                ip.paragraph_format.space_before = Pt(0)
                ip.paragraph_format.space_after = Pt(1)
                w_run = ip.add_run(w_kanji)
                w_run.font.name = "MS Mincho"
                w_run.font.size = Pt(17)
                w_run.font.bold = True

                hp = info_cell.add_paragraph()
                hp.paragraph_format.space_before = Pt(0)
                hp.paragraph_format.space_after = Pt(1)
                h_run = hp.add_run(w_hira)
                h_run.font.name = "Hiragino Mincho ProN"
                h_run.font.size = Pt(8.5)
                h_run.font.color.rgb = RGBColor(0x8C, 0x6B, 0x3E)

                ep = info_cell.add_paragraph()
                ep.paragraph_format.space_before = Pt(0)
                ep.paragraph_format.space_after = Pt(0)
                e_run = ep.add_run(w_eng)
                e_run.font.name = "Arial"
                e_run.font.size = Pt(7)
                e_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

            # Spacing between Kanji blocks
            sp = doc.add_paragraph()
            sp.paragraph_format.space_before = Pt(0)
            sp.paragraph_format.space_after = Pt(10)

        # Page break at the end of each Set
        if set_idx < len(subset_dict) - 1:
            doc.add_page_break()

    doc.save(filename)
    print(f"Successfully generated: {filename}")

if __name__ == "__main__":
    # Generate complete Sets 1 to 16
    generate_kanji_document(DATA, "Kanji_Practice_Sheets_Sets_1_to_16.docx")
    
    # Generate Sets 1 to 10 for convenience
    sets_1_10_sub = {f"Set {i}": DATA[f"Set {i}"] for i in range(1, 11) if f"Set {i}" in DATA}
    generate_kanji_document(sets_1_10_sub, "Kanji_Practice_Sheets_Sets_1_to_10.docx")
