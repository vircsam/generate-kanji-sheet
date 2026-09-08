import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from soma import (
    CHAPTER_1_DATA,
    CHAPTER_2_DATA,
    CHAPTER_3_DATA,
    CHAPTER_4_DATA,
    CHAPTER_5_DATA,
    CHAPTER_6_DATA
)

from generate_pdf_skanji import create_skanji_pdf

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="D3C7B6"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'<w:right w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def generate_skanji_document(subset_dict, filename="Kanji_Practice_3cols.docx"):
    doc = Document()
    
    # Configure margins for 7.7 inches printable width
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.4)
        s.bottom_margin = Inches(0.4)
        s.left_margin = Inches(0.4)
        s.right_margin = Inches(0.4)

    # 3 Columns: 1 KANJI+MEANING (2.7 in) + 2 Practice columns (2.5 in each) = 7.7 in total
    COL_WIDTHS = [Inches(2.7), Inches(2.5), Inches(2.5)]
    
    for set_idx, (set_name, kanji_list) in enumerate(subset_dict.items()):
        p_set = doc.add_paragraph()
        p_set.paragraph_format.space_before = Pt(6)
        p_set.paragraph_format.space_after = Pt(4)
        run_set = p_set.add_run(set_name)
        run_set.font.name = "Georgia"
        run_set.font.size = Pt(16)
        run_set.font.bold = True
        run_set.font.color.rgb = RGBColor(0x8C, 0x2A, 0x1E)

        # 3-Column Table
        table = doc.add_table(rows=1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table, color="D3C7B6")

        # Header Row
        hdr_row = table.rows[0]
        hdr_titles = ["KANJI & MEANING", "PRACTICE 1", "PRACTICE 2"]
        for i, cell in enumerate(hdr_row.cells):
            cell.width = COL_WIDTHS[i]
            set_cell_background(cell, "EDE5D8")
            set_cell_margins(cell, top=40, bottom=40, left=60, right=60)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(hdr_titles[i])
            r.font.name = "Arial"
            r.font.size = Pt(9)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x77, 0x77, 0x77)

        # Kanji rows (No vocabulary, only Kanji + Meaning/Readings in Column 0, Columns 1 & 2 empty)
        for kanji, meaning, _vocab_list in kanji_list:
            row = table.add_row()
            trPr = row._element.get_or_add_trPr()
            trHeight = parse_xml(f'<w:trHeight {nsdecls("w")} w:val="850" w:hRule="atLeast"/>')
            trPr.append(trHeight)

            # Set column widths & vertical alignment
            for i, c in enumerate(row.cells):
                c.width = COL_WIDTHS[i]
                c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

            # Column 0: Kanji + Meaning
            c0 = row.cells[0]
            set_cell_background(c0, "F9F6F0")
            set_cell_margins(c0, top=50, bottom=50, left=80, right=60)
            
            p_kanji = c0.paragraphs[0]
            p_kanji.paragraph_format.space_before = Pt(1)
            p_kanji.paragraph_format.space_after = Pt(1)
            r_kanji = p_kanji.add_run(kanji)
            r_kanji.font.name = "MS Mincho"
            r_kanji.font.size = Pt(22)
            r_kanji.font.bold = True

            p_meaning = c0.add_paragraph()
            p_meaning.paragraph_format.space_before = Pt(0)
            p_meaning.paragraph_format.space_after = Pt(1)
            r_meaning = p_meaning.add_run(meaning)
            r_meaning.font.name = "Georgia"
            r_meaning.font.size = Pt(9.5)
            r_meaning.font.italic = True
            r_meaning.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

            # Column 1 & 2: Empty practice cells
            for col_idx in [1, 2]:
                c_empty = row.cells[col_idx]
                set_cell_margins(c_empty, top=50, bottom=50, left=60, right=60)

        # Page break between sets if multiple sets in subset_dict
        if set_idx < len(subset_dict) - 1:
            doc.add_page_break()

    doc.save(filename)
    print(f"Successfully generated 3-column Kanji sheet: {filename}")

if __name__ == "__main__":
    # Combined dataset for Chapters 1-6
    all_chapters = {}
    for ch in [CHAPTER_1_DATA, CHAPTER_2_DATA, CHAPTER_3_DATA, CHAPTER_4_DATA, CHAPTER_5_DATA, CHAPTER_6_DATA]:
        all_chapters.update(ch)

    # 1. Generate Combined 3-Column DOCX for Chapters 1-6
    generate_skanji_document(all_chapters, "Kanji_Practice_All_Chapters_3cols.docx")

    # 2. Generate Individual Chapter 3-Column DOCX files
    generate_skanji_document(CHAPTER_1_DATA, "Kanji_Practice_Chapter1_3cols.docx")
    generate_skanji_document(CHAPTER_2_DATA, "Kanji_Practice_Chapter2_3cols.docx")
    generate_skanji_document(CHAPTER_3_DATA, "Kanji_Practice_Chapter3_3cols.docx")
    generate_skanji_document(CHAPTER_4_DATA, "Kanji_Practice_Chapter4_3cols.docx")
    generate_skanji_document(CHAPTER_5_DATA, "Kanji_Practice_Chapter5_3cols.docx")
    generate_skanji_document(CHAPTER_6_DATA, "Kanji_Practice_Chapter6_3cols.docx")

    # 3. Generate Combined 3-Column PDF for skanji
    create_skanji_pdf(all_chapters, "Kanji_Practice_All_Chapters_3cols.pdf")

